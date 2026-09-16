import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Tuple

TEMPLATE_VERSION = "thpay_admissoes_v1.0"

COLUMNS = [
    ("Nome Completo", "Texto obrigatório", "Ex: Ana Clara dos Santos"),
    ("CPF", "11 dígitos (com ou sem pontuação)", "Ex: 123.456.789-00"),
    ("Data de Nascimento", "Formato DD/MM/AAAA", "Ex: 15/04/1995"),
    ("E-mail", "E-mail para convite de pré-admissão", "Ex: ana.santos@email.com"),
    ("Celular", "DDD + Número", "Ex: (11) 98765-4321"),
    ("CEP", "8 dígitos", "Ex: 01310-100"),
    ("Logradouro", "Rua / Avenida", "Ex: Avenida Paulista"),
    ("Número", "Número residencial", "Ex: 1000"),
    ("Complemento", "Apto, Bloco (opcional)", "Ex: Apto 42"),
    ("Bairro", "Bairro residencial", "Ex: Bela Vista"),
    ("Cidade", "Município", "Ex: São Paulo"),
    ("UF", "Sigla do estado (2 letras)", "Ex: SP"),
    ("Cargo", "Título da função contratual", "Ex: Engenheiro de Software Júnior"),
    ("CBO", "Código CBO 2002 (opcional)", "Ex: 2124-05"),
    ("Salário Base", "Valor nominal em R$", "Ex: 5500,00"),
    ("Departamento", "Nome da área interna", "Ex: Engenharia de Software"),
    ("Centro de Custo", "Código ou nome", "Ex: CC-1001"),
    ("Filial", "Unidade contratante", "Ex: Matriz São Paulo"),
    ("Data de Admissão", "Data prevista de início (DD/MM/AAAA)", "Ex: 01/10/2026"),
    ("Modalidade", "CLT, Estágio, Aprendiz, Temporário", "Ex: CLT"),
    ("Jornada", "Carga horária semanal", "Ex: 40h semanais"),
    ("Banco", "Nome ou código da instituição", "Ex: Banco do Brasil"),
    ("Agência", "Número da agência", "Ex: 1234"),
    ("Conta Corrente", "Número da conta com dígito", "Ex: 56789-0"),
    ("Chave Pix", "Tipo e chave (opcional)", "Ex: 123.456.789-00"),
    ("Opção VT", "Sim ou Não", "Ex: Sim"),
    ("Divisão VA/VR", "Percentual VA (o restante será VR)", "Ex: 50")
]

class AdmissionTemplateGenerator:
    @staticmethod
    def generate_xlsx() -> bytes:
        wb = openpyxl.Workbook()
        
        # 1. Aba COLABORADORES
        ws_main = wb.active
        ws_main.title = "COLABORADORES"
        ws_main.views.sheetView[0].showGridLines = True
        
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="101A44", end_color="101A44", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='E5DED5'),
            right=Side(style='thin', color='E5DED5'),
            top=Side(style='thin', color='E5DED5'),
            bottom=Side(style='thin', color='E5DED5')
        )
        
        for col_idx, (col_name, _, _) in enumerate(COLUMNS, start=1):
            cell = ws_main.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            ws_main.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = max(len(col_name) + 5, 14)
            
        # 2. Aba DICIONARIO_DE_CAMPOS
        ws_dict = wb.create_sheet(title="DICIONARIO_DE_CAMPOS")
        ws_dict.append(["Campo", "Instruções de Preenchimento", "Exemplo Válido"])
        for cell in ws_dict[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F66F3", end_color="2F66F3", fill_type="solid")
        for col_name, rule, example in COLUMNS:
            ws_dict.append([col_name, rule, example])
        ws_dict.column_dimensions["A"].width = 25
        ws_dict.column_dimensions["B"].width = 45
        ws_dict.column_dimensions["C"].width = 30

        # 3. Aba VALORES_ACEITOS
        ws_values = wb.create_sheet(title="VALORES_ACEITOS")
        ws_values.append(["Tipo de Dado", "Valores Permitidos / Padrões do Sistema"])
        for cell in ws_values[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F66F3", end_color="2F66F3", fill_type="solid")
        ws_values.append(["Modalidade Contratual", "CLT, Aprendiz, Estágio, Temporário, Intermitente"])
        ws_values.append(["Opção VT", "Sim (solicita benefício), Não (termo de renúncia / opt-out)"])
        ws_values.append(["Divisão VA/VR", "0 a 100 (ex: 50 indica 50% VA e 50% VR; soma sempre 100%)"])
        ws_values.append(["Formato de Data", "DD/MM/AAAA ou AAAA-MM-DD"])
        ws_values.append(["Estado / UF", "SP, RJ, MG, RS, PR, SC, BA, PE, CE, etc. (2 letras maiúsculas)"])
        ws_values.column_dimensions["A"].width = 25
        ws_values.column_dimensions["B"].width = 60

        # 4. Aba EXEMPLO_PREENCHIDO
        ws_example = wb.create_sheet(title="EXEMPLO_PREENCHIDO")
        ws_example.append([col[0] for col in COLUMNS])
        for cell in ws_example[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws_example.append([col[2].replace("Ex: ", "") for col in COLUMNS])
        for col_idx in range(1, len(COLUMNS) + 1):
            ws_example.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def generate_csv() -> str:
        headers = [col[0] for col in COLUMNS]
        example = [col[2].replace("Ex: ", "") for col in COLUMNS]
        import csv
        out = io.StringIO()
        writer = csv.writer(out, delimiter=";")
        writer.writerow(headers)
        writer.writerow(example)
        return out.getvalue()
