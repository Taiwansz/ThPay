import csv
import io
from typing import Any, Dict, List, Optional, Tuple

class SpreadsheetParser:
    """
    Parser robusto para arquivos XLSX e CSV (com deteccao automatica de delimitador e encoding).
    """
    @staticmethod
    def parse_csv(content_bytes: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
        # Tentar detectar encoding
        text = None
        for enc in ["utf-8-sig", "utf-8", "latin-1", "cp1252"]:
            try:
                text = content_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
                
        if text is None:
            text = content_bytes.decode("utf-8", errors="replace")

        # Detectar delimitador analisando as primeiras linhas
        sample = text[:4096]
        sniffer = csv.Sniffer()
        try:
            dialect = sniffer.sniff(sample, delimiters=[",", ";", "\t", "|"])
            delimiter = dialect.delimiter
        except Exception:
            # Fallback heuristic
            if sample.count(";") > sample.count(","):
                delimiter = ";"
            elif sample.count("\t") > sample.count(","):
                delimiter = "\t"
            else:
                delimiter = ","

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        rows_raw = [r for r in reader if any(cell.strip() for cell in r)]
        if not rows_raw:
            return [], []

        headers = [h.strip() for h in rows_raw[0]]
        data_rows: List[Dict[str, Any]] = []
        
        for idx, row in enumerate(rows_raw[1:], start=2):
            row_dict = {}
            for col_idx, header in enumerate(headers):
                val = row[col_idx].strip() if col_idx < len(row) else ""
                row_dict[header] = val
            data_rows.append(row_dict)

        return headers, data_rows

    @staticmethod
    def parse_xlsx(content_bytes: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content_bytes), data_only=True)
        
        # Preferir aba COLABORADORES se existir, senao primeira aba ativa
        sheet = None
        for sheet_name in wb.sheetnames:
            if sheet_name.upper().strip() in ("COLABORADORES", "ADMISSOES", "FUNCIONARIOS"):
                sheet = wb[sheet_name]
                break
        if sheet is None:
            sheet = wb.active

        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            return [], []

        headers = [str(h).strip() if h is not None else f"Coluna_{idx+1}" for idx, h in enumerate(header_row)]
        
        data_rows: List[Dict[str, Any]] = []
        for row in rows_iter:
            if not any(val is not None and str(val).strip() != "" for val in row):
                continue
            row_dict = {}
            for col_idx, header in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                if val is None:
                    row_dict[header] = ""
                elif isinstance(val, (int, float)):
                    row_dict[header] = str(val)
                else:
                    row_dict[header] = str(val).strip()
            data_rows.append(row_dict)

        return headers, data_rows

    @classmethod
    def parse(cls, filename: str, content_bytes: bytes) -> Tuple[List[str], List[Dict[str, Any]]]:
        lower_name = filename.lower()
        if lower_name.endswith(".xlsx") or lower_name.endswith(".xls"):
            return cls.parse_xlsx(content_bytes)
        else:
            return cls.parse_csv(content_bytes)
